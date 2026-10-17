# -*- coding: utf-8 -*-
"""postapply_check.py —— 写回后确定性终检一条龙（替代每轮手抄的 verify_output /
openxml_check / check_elements / integrity 四件套）。

写回（apply_track_changes.py）之后必须机械核一遍"修订到底写对了没有"。历史上 12 个
运行目录各写一份终检脚本，口径还各不相同；本脚本把**能确定判定的**全部固化：

  1. OpenXML 包体完整：所有 .xml/.rels 部件可解析；w:ins/w:del 都带 id+author；
  2. 段落总数：产出 docx 与源 docx 的正文顶层 w:p 数一致；
  3. accept 视图（全部 w:t，不含 del 里的）逐段 == 期望 newText；
  4. reject 视图（非 ins 的 w:t + delText）逐段 == 期望原文（original）——
     accept/reject 两个视图同时成立，才说明修订是"可逐条接受/拒绝"的真修订；
  5. 修订段数：产出里 accept≠reject 的段数 == 期望里 newText≠原文 的段数，
     且**段 id 集合逐一对应**（少写/多写都会指名道姓）；
  6. 批注：comments.xml 条数 == 期望 flag 条数；commentRangeStart/End/commentReference
     与 comments.xml 的 id 集合对齐；
  7. run 级附加元素前后计数：w:tab / w:drawing / w:fldChar / w:bookmarkStart / m:oMath
     **总数与逐段计数都必须一致**——历史事故：重建段落时物理删掉过 OMML 公式与 _Toc 书签。

用法:
  python postapply_check.py <产出docx> <merged.json> [--src 源docx] [-o report.json] [--quiet]
  例: python postapply_check.py "6-xxx-校正修订.docx" final.json --src "src.docx" -o check.json

参数:
  <merged.json>  写回时喂给 apply_track_changes.py 的那份 json（**收口后请传 final.json**；
                 若还没跑 finalize.py 就传 merged.json）。每条 {id,paraIndex,original,newText,flag?}
  --src          写回前的源 docx（不给则跳过"段落总数/附加元素前后一致"两项，并告警）
  -o             机器可读报告，默认 <产出docx 同目录>/<stem>.postapply-check.json
  --quiet        只打结论行

退出码: 0 全部通过；1 有硬失败（控制台指明是哪一项，报告里给出 id 明细）。
"""
import argparse
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
# run 级 / 段落级"附加元素"：重建段落时最容易被静默丢掉的几类
EXTRA_TAGS = [
    ('w:tab', W + 'tab'),
    ('w:drawing', W + 'drawing'),
    ('w:fldChar', W + 'fldChar'),
    ('w:bookmarkStart', W + 'bookmarkStart'),
    ('m:oMath', M + 'oMath'),
]


def load_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save_json(path, obj):
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=1), encoding='utf-8')


def body_paras(doc):
    """正文顶层段落（与 extract_body_paragraphs.py 的 paraIndex 口径一致）。"""
    return [c for c in doc.element.body if c.tag == W + 'p']


def view_text(p, accept):
    """accept=True：全部 w:t（不含 del 里的、不含 delText）；
    accept=False：非 ins 的 w:t + delText（即"拒绝全部修订"后的文本）。"""
    out = []
    for t in p.iter():
        if not isinstance(t.tag, str):
            continue
        if t.tag not in (W + 't', W + 'delText'):
            continue
        is_deltext = (t.tag == W + 'delText')
        in_ins = in_del = False
        node = t
        while node is not None:
            if isinstance(node.tag, str):
                ln = etree.QName(node).localname
                if ln == 'ins':
                    in_ins = True
                elif ln == 'del':
                    in_del = True
            node = node.getparent()
        if accept:
            if not is_deltext and not in_del:
                out.append(t.text or '')
        else:
            if is_deltext or in_del:
                out.append(t.text or '')
            elif not in_ins:
                out.append(t.text or '')
    return ''.join(out)


def extras_sig(paras):
    """返回 (总数 Counter, 逐段 [{tag: n}])。"""
    total = Counter()
    per = []
    for p in paras:
        row = {}
        for name, tag in EXTRA_TAGS:
            n = len(p.findall('.//' + tag))
            row[name] = n
            total[name] += n
        per.append(row)
    return total, per


def package_sanity(docx_path):
    """zip 内所有 xml/rels 可解析 + 修订属性完备。"""
    detail = {'parts': 0, 'malformed': [], 'ins': 0, 'del': 0, 'bad_attrs': [],
              'comment_refs': 0, 'comments': -1, 'comment_ids_aligned': False,
              'range_start': 0, 'range_end': 0}
    try:
        z = zipfile.ZipFile(str(docx_path))
    except Exception as e:
        detail['malformed'].append(('__zip__', str(e)[:120]))
        return detail
    names = z.namelist()
    detail['parts'] = len(names)
    for n in names:
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                etree.fromstring(z.read(n))
            except Exception as e:
                detail['malformed'].append((n, str(e)[:120]))
    if 'word/document.xml' in names:
        doc = etree.fromstring(z.read('word/document.xml'))
        ins = doc.findall('.//' + W + 'ins')
        dele = doc.findall('.//' + W + 'del')
        detail['ins'], detail['del'] = len(ins), len(dele)
        for el in ins + dele:
            if el.get(W + 'id') is None or not el.get(W + 'author'):
                detail['bad_attrs'].append((etree.QName(el).localname,
                                            el.get(W + 'id'), el.get(W + 'author')))
        refs = doc.findall('.//' + W + 'commentReference')
        detail['comment_refs'] = len(refs)
        detail['range_start'] = len(doc.findall('.//' + W + 'commentRangeStart'))
        detail['range_end'] = len(doc.findall('.//' + W + 'commentRangeEnd'))
        ref_ids = sorted(c.get(W + 'id') for c in refs)
        if 'word/comments.xml' in names:
            com = etree.fromstring(z.read('word/comments.xml'))
            cs = com.findall('.//' + W + 'comment')
            detail['comments'] = len(cs)
            com_ids = sorted(c.get(W + 'id') for c in cs)
            detail['comment_ids_aligned'] = (ref_ids == com_ids)
        else:
            detail['comment_ids_aligned'] = (not ref_ids)
    return detail


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument('docx')
    ap.add_argument('merged')
    ap.add_argument('--src', default='', help='写回前的源 docx（用于段落数/附加元素前后比对）')
    ap.add_argument('--summary', default='',
                    help='apply_track_changes 的 *.summary.json（默认自动找 <产出docx主干>.summary.json）。'
                         '用于把「保护清单/含特殊内容而有意跳过」的段从期望中排除，避免误判失败')
    ap.add_argument('-o', '--out', default='')
    ap.add_argument('--quiet', action='store_true')
    args = ap.parse_args()

    out_docx = Path(args.docx).resolve()
    merged_path = Path(args.merged).resolve()
    src_docx = Path(args.src).resolve() if args.src else None
    if not out_docx.exists():
        print(f'✗ 产出 docx 不存在: {out_docx}')
        sys.exit(1)
    if not merged_path.exists():
        print(f'✗ merged/final json 不存在: {merged_path}')
        sys.exit(1)

    entries = load_json(merged_path)
    if not isinstance(entries, list):
        print(f'✗ {merged_path.name} 应为数组（rewrites 格式）')
        sys.exit(1)
    checks, failures, warnings = [], [], []

    def add(name, ok, detail, hard=True):
        checks.append({'name': name, 'ok': bool(ok), 'hard': bool(hard),
                       'detail': detail})
        if not ok:
            (failures if hard else warnings).append(name)
        if not args.quiet or not ok:
            mark = '✓' if ok else ('✗' if hard else '!')
            print(f'{mark} {name}: {detail}')

    print(f'产出: {out_docx.name}')
    print(f'期望: {merged_path.name}（{len(entries)} 条）')
    print(f'源文件: {src_docx.name if src_docx else "(未提供)"}')
    print('-' * 60)

    # ---- 1. OpenXML 包体 ----
    pkg = package_sanity(out_docx)
    ok = (not pkg['malformed'] and not pkg['bad_attrs']
          and pkg['comment_ids_aligned'] and pkg['range_start'] == pkg['comments']
          and pkg['range_end'] == pkg['comments'])
    add('openxml_package',
        ok,
        f'部件 {pkg["parts"]}，坏部件 {len(pkg["malformed"])}，w:ins {pkg["ins"]}，'
        f'w:del {pkg["del"]}，属性缺失 {len(pkg["bad_attrs"])}，'
        f'批注定义 {pkg["comments"]} / 引用 {pkg["comment_refs"]} / 锚点 '
        f'{pkg["range_start"]}→{pkg["range_end"]}，id 对齐={pkg["comment_ids_aligned"]}'
        + (f'；坏部件 {pkg["malformed"][:3]}' if pkg['malformed'] else ''))
    if pkg['malformed']:
        warnings.append('malformed_parts')

    # ---- 打开 docx ----
    try:
        doc = Document(str(out_docx))
    except Exception as e:
        add('open_docx', False, f'python-docx 打不开: {e}')
        save_json(Path(args.out) if args.out else out_docx.with_name(
            out_docx.stem + '.postapply-check.json'),
            {'tool': 'postapply_check.py', 'ok': False, 'checks': checks,
             'failures': failures, 'warnings': warnings})
        print('\n结论: 产出 docx 打不开，终检中止')
        sys.exit(1)
    out_paras = body_paras(doc)
    add('open_docx', True, f'正文顶层段落 {len(out_paras)}')

    # ---- 2. 段落总数（需 --src）----
    if src_docx and src_docx.exists():
        src_doc = Document(str(src_docx))
        src_paras = body_paras(src_doc)
        add('paragraph_count', len(src_paras) == len(out_paras),
            f'源 {len(src_paras)} → 产出 {len(out_paras)}'
            + ('' if len(src_paras) == len(out_paras) else '（段落总数变了！）'))
    else:
        add('paragraph_count', True, '未提供 --src，跳过（建议补上）', hard=False)
        warnings.append('paragraph_count_skipped')

    # ---- 3/4. accept / reject 双视图逐段比对 ----
    # 写回阶段会有意跳过两类段（保护清单 / 段落含特殊内容），它们的 accept 视图必然仍是原文；
    # 若不排除，会被误判成"期望有但未写"。优先从 apply 的 summary.json 里取这份跳过名单。
    skip_ids = set()
    summ_path = Path(args.summary) if args.summary else \
        out_docx.with_name(out_docx.stem + '.summary.json')
    if summ_path.exists():
        try:
            sm = json.loads(summ_path.read_text(encoding='utf-8'))
            for r in (sm.get('skipReasons') or []):
                head = str(r).split()[0] if str(r).split() else ''
                if head.startswith('id') and head[2:].isdigit():
                    skip_ids.add(int(head[2:]))
            add('apply_summary', True,
                f'{summ_path.name}：写回有意跳过 {len(skip_ids)} 段（已从期望中排除）', hard=False)
        except Exception as e:
            warnings.append('apply_summary_unreadable')
            add('apply_summary', True, f'{summ_path.name} 读取失败，忽略: {e}', hard=False)
    else:
        warnings.append('apply_summary_missing')
        add('apply_summary', True,
            f'未找到 {summ_path.name}；若用了保护清单，accept_view/revision_count 可能误报（可加 --summary）',
            hard=False)

    acc_bad, rej_bad, exp_changed, act_changed = [], [], [], []
    idx_seen = {}
    for e in entries:
        pid, idx = e.get('id'), e.get('paraIndex')
        o = e.get('original')
        nt = e.get('newText')
        if not isinstance(pid, int) or not isinstance(idx, int):
            acc_bad.append((pid, 'entry 缺 id/paraIndex'))
            continue
        if idx in idx_seen:
            acc_bad.append((pid, f'paraIndex {idx} 与 id{idx_seen[idx]} 重复'))
            continue
        idx_seen[idx] = pid
        if pid in skip_ids:
            continue                        # 写回阶段有意跳过：不参与双视图与修订计数
        if not (0 <= idx < len(out_paras)):
            acc_bad.append((pid, f'paraIndex {idx} 越界'))
            continue
        p = out_paras[idx]
        av = view_text(p, True).strip()
        rv = view_text(p, False).strip()
        exp_nt = (nt or '').strip()
        exp_o = (o or '').strip()
        if av != exp_nt:
            acc_bad.append((pid, f'accept≠newText exp={exp_nt[:40]!r} got={av[:40]!r}'))
        if rv != exp_o:
            rej_bad.append((pid, f'reject≠original exp={exp_o[:40]!r} got={rv[:40]!r}'))
        if exp_nt != exp_o:
            exp_changed.append(pid)
            if av != rv:
                act_changed.append(pid)
    add('accept_view', not acc_bad,
        f'不匹配 {len(acc_bad)} 段' + (f'，前 6: {[x[0] for x in acc_bad[:6]]}' if acc_bad else ''))
    if acc_bad:
        print('    例: ' + '; '.join(str(x[1]) for x in acc_bad[:3]))
    add('reject_view', not rej_bad,
        f'不匹配 {len(rej_bad)} 段' + (f'，前 6: {[x[0] for x in rej_bad[:6]]}' if rej_bad else ''))
    if rej_bad:
        print('    例: ' + '; '.join(str(x[1]) for x in rej_bad[:3]))

    # ---- 5. 修订段数 ----
    actual_all_changed = [i for i, p in enumerate(out_paras)
                          if view_text(p, True).strip() != view_text(p, False).strip()]
    exp_set, act_set = set(exp_changed), set(idx_seen[i] for i in actual_all_changed)
    only_exp = sorted(exp_set - act_set)
    only_act = sorted(act_set - exp_set)
    add('revision_count',
        len(exp_set) == len(actual_all_changed) and not only_exp and not only_act,
        f'期望修订 {len(exp_set)} 段，实际修订 {len(actual_all_changed)} 段'
        + (f'；期望有但未写: {only_exp[:8]}' if only_exp else '')
        + (f'；写了但期望没有: {only_act[:8]}' if only_act else ''))

    # ---- 6. 批注 ----
    exp_flags = [e.get('id') for e in entries if str(e.get('flag') or '').strip()]
    n_com = pkg['comments']
    add('comments_count', n_com == len(exp_flags),
        f'期望 flag {len(exp_flags)} 条，产出批注 {n_com} 条'
        + ('' if n_com == len(exp_flags) else f'；差额 {len(exp_flags) - n_com}'))

    # ---- 7. run 级附加元素前后一致 ----
    if src_docx and src_docx.exists():
        src_paras = body_paras(Document(str(src_docx)))
        tot_s, per_s = extras_sig(src_paras)
        tot_o, per_o = extras_sig(out_paras)
        tot_diff = {k: (tot_s[k], tot_o[k]) for k in tot_s if tot_s[k] != tot_o[k]}
        n = min(len(per_s), len(per_o))
        per_diff = []
        for i in range(n):
            for k in per_s[i]:
                if per_s[i][k] != per_o[i][k]:
                    per_diff.append((i, k, per_s[i][k], per_o[i][k]))
        add('extras_count', not tot_diff and not per_diff,
            '总数 ' + ' '.join(f'{k}={tot_o[k]}' for k, _ in EXTRA_TAGS) + '；'
            + ('前后一致' if not tot_diff else f'总数变化 {tot_diff}')
            + (f'；逐段变化 {len(per_diff)} 处（例 {per_diff[:3]}）' if per_diff else '；逐段一致'))
    else:
        add('extras_count', True, '未提供 --src，跳过（历史事故：OMML/书签被静默删除）', hard=False)
        warnings.append('extras_count_skipped')

    # ---- 汇总 ----
    report = {
        'tool': 'postapply_check.py',
        'version': 1,
        'ok': not failures,
        'docx': str(out_docx),
        'merged': str(merged_path),
        'src': str(src_docx) if src_docx else '',
        'expected_changed': len(exp_changed),
        'actual_changed': len(actual_all_changed),
        'expected_flags': len(exp_flags),
        'comments': n_com,
        'accept_mismatch': [{'id': p, 'detail': d} for p, d in acc_bad],
        'reject_mismatch': [{'id': p, 'detail': d} for p, d in rej_bad],
        'changed_missing': only_exp,
        'changed_extra': only_act,
        'checks': checks,
        'failures': failures,
        'warnings': warnings,
    }
    out = Path(args.out) if args.out else out_docx.with_name(
        out_docx.stem + '.postapply-check.json')
    save_json(out, report)

    print('-' * 60)
    if failures:
        print(f'✗ 终检失败：{failures}')
        if 'revision_count' in failures or 'accept_view' in failures:
            print('  提示：若你传的是 merged.json 而写回用的是收口后的 final.json，'
                  '两者条数会不一致——请传写回时实际使用的那份。')
        print(f'报告 → {out}')
        sys.exit(1)
    print(f'✓ 终检通过：修订 {len(actual_all_changed)} 段 / 批注 {n_com} 条 / '
          f'双视图一致 / 附加元素一致')
    print(f'报告 → {out}')
    return 0


if __name__ == '__main__':
    sys.exit(main() or 0)
